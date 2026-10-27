import base64
import io
import json
import os
import re

import streamlit as st
from docx import Document
from openai import OpenAI

st.set_page_config(page_title="AI Exam Paper Checker", page_icon="📝", layout="wide")

OPENROUTER_BASE_URL = "https://openrouter.ai/api/v1"
DEFAULT_MODEL = "openrouter/free"  # random free model; use one specific model slug for consistent grading

PROMPT = """You are a fair and consistent exam paper examiner.
Check the STUDENT PAPER below. It may contain only answers listed by question number, without the question text.
In that case, match each answer to its question using the QUESTION PAPER or the MARKING SCHEME / ANSWER KEY.
If a MARKING SCHEME / ANSWER KEY is provided, grade strictly against it.
If none is provided, use your own subject knowledge to judge the correctness of each answer.

Rules:
- Identify every question in the paper and award marks for each one.
- If the paper shows marks per question, use them. Otherwise split TOTAL_MARKS sensibly across the questions.
- Never award more than the maximum for a question. Unanswered questions get 0.
- Be objective. Give a short reason for every mark you award or deduct.
- If the paper is unreadable or is not an exam paper, say so in "summary" and give 0 marks.
{EXTRA}

Return ONLY valid JSON in exactly this shape:
{
  "student_name": "name if written on the paper, else Unknown",
  "subject": "subject if identifiable, else Unknown",
  "questions": [
    {"number": "Q1", "max_marks": 10, "awarded_marks": 7, "feedback": "short reason"}
  ],
  "strengths": ["..."],
  "weaknesses": ["..."],
  "summary": "2-3 sentence overall assessment"
}

TOTAL_MARKS: {TOTAL}
"""


def get_api_key() -> str | None:
 
    try:
        key = st.secrets.get("OPENROUTER_API_KEY")
    except Exception:
        key = None
    return key or os.getenv("OPENROUTER_API_KEY")


def docx_to_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def to_parts(uploaded, label: str) -> list:
   
    data = uploaded.getvalue()
    name = uploaded.name.lower()
    if name.endswith(".pdf"):
        b64 = base64.b64encode(data).decode("utf-8")
        return [
            {"type": "text", "text": f"{label}:"},
            {"type": "file", "file": {"filename": uploaded.name, "file_data": f"data:application/pdf;base64,{b64}"}},
        ]
    if name.endswith(".docx"):
        return [{"type": "text", "text": f"{label}:\n{docx_to_text(data)}"}]
    return [{"type": "text", "text": f"{label}:\n{data.decode('utf-8', errors='ignore')}"}]


def parse_json(text: str) -> dict:
   
    text = re.sub(r"```(?:json)?", "", text or "").strip()
    start, end = text.find("{"), text.rfind("}")
    if start == -1 or end == -1:
        raise json.JSONDecodeError("No JSON found", text, 0)
    return json.loads(text[start : end + 1])


def evaluate(api_key, model, paper, scheme, total_marks, extra_rules, questions=None) -> dict:
    client = OpenAI(base_url=OPENROUTER_BASE_URL, api_key=api_key)
    extra = f"- Extra instructions from the teacher: {extra_rules}" if extra_rules else ""
    prompt = PROMPT.replace("{TOTAL}", str(total_marks)).replace("{EXTRA}", extra)

    contents = [{"type": "text", "text": prompt}] + to_parts(paper, "STUDENT PAPER")
    if questions:
        contents += to_parts(questions, "QUESTION PAPER")
    if scheme:
        contents += to_parts(scheme, "MARKING SCHEME / ANSWER KEY")

    response = client.chat.completions.create(
        model=model,
        messages=[{"role": "user", "content": contents}],
        temperature=0.2,
    )
    return parse_json(response.choices[0].message.content)


def grade_letter(pct: float) -> str:
    for cutoff, letter in [(90, "A+"), (80, "A"), (70, "B"), (60, "C"), (50, "D")]:
        if pct >= cutoff:
            return letter
    return "F"


# ---------------- UI ----------------
st.title("📝 AI Exam Paper Checker")
st.caption("Upload a student's exam paper and get a score with feedback, powered by OpenRouter.")

with st.sidebar:
    st.header("Settings")
    api_key = get_api_key() or st.text_input("OpenRouter API key", type="password")
    model_name = st.text_input("OpenRouter model", value=DEFAULT_MODEL)
    total_marks = st.number_input("Total marks", min_value=1, value=100, step=1)
    extra_rules = st.text_area("Extra grading instructions (optional)", placeholder="e.g. Be lenient on spelling.")
    st.info("AI scores can be wrong. Review the result before using it as a final grade.")

col1, col2, col3 = st.columns(3)
with col1:
    paper = st.file_uploader("Student paper / answer sheet (required)", type=["pdf", "docx"])
with col2:
    questions = st.file_uploader("Question paper (optional)", type=["pdf", "docx", "txt"])
with col3:
    scheme = st.file_uploader("Marking scheme / answer key (optional)", type=["pdf", "docx", "txt"])

if st.button("Check paper", type="primary", disabled=paper is None):
    if not api_key:
        st.error("Please add your OpenRouter API key in the sidebar.")
        st.stop()

    with st.spinner("Checking the paper..."):
        try:
            result = evaluate(api_key, model_name, paper, scheme, total_marks, extra_rules, questions)
        except json.JSONDecodeError:
            st.error("The AI returned an unreadable response. Please try again.")
            st.stop()
        except Exception as e:
            st.error(f"Something went wrong: {e}")
            st.stop()

    questions = result.get("questions", [])
    score = sum(float(q.get("awarded_marks", 0)) for q in questions)
    max_score = sum(float(q.get("max_marks", 0)) for q in questions) or float(total_marks)
    score = min(score, max_score)
    pct = (score / max_score * 100) if max_score else 0

    st.subheader(f"Result: {result.get('student_name', 'Unknown')} - {result.get('subject', 'Unknown')}")
    m1, m2, m3 = st.columns(3)
    m1.metric("Score", f"{score:g} / {max_score:g}")
    m2.metric("Percentage", f"{pct:.1f}%")
    m3.metric("Grade", grade_letter(pct))
    st.progress(min(int(pct), 100))

    st.write(result.get("summary", ""))

    if questions:
        st.markdown("#### Question-wise marks")
        st.dataframe(
            [
                {
                    "Question": q.get("number", ""),
                    "Awarded": q.get("awarded_marks", 0),
                    "Max": q.get("max_marks", 0),
                    "Feedback": q.get("feedback", ""),
                }
                for q in questions
            ],
            use_container_width=True,
            hide_index=True,
        )

    c1, c2 = st.columns(2)
    with c1:
        st.markdown("#### Strengths")
        for s in result.get("strengths", []):
            st.write(f"- {s}")
    with c2:
        st.markdown("#### Weaknesses")
        for w in result.get("weaknesses", []):
            st.write(f"- {w}")

    st.download_button(
        "Download result (JSON)",
        data=json.dumps(result, indent=2),
        file_name="result.json",
        mime="application/json",
    )
